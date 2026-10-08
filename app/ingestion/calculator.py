"""Import des référentiels depuis le calculateur Excel (ex. « B80 105.xlsm »).

Utilisé par l'écran Référentiels et par `tools/extract_masters.py`. Les CSV produits
restent dans le dossier local des données (jamais dans le dépôt). La structure n'est
pas supposée : pour chaque onglet, l'en-tête est cherché parmi plusieurs noms possibles.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.ingestion.tables import (
    TableSpec,
    ValidationReport,
    parse_dates,
    parse_table,
    read_raw,
)


@dataclass(frozen=True)
class SheetJob:
    sheet_names: tuple[str, ...]
    spec: TableSpec
    output: str
    out_columns: tuple[str, ...]
    date_fields: tuple[str, ...] = ()
    fill_boutique: bool = False  # onglet propre à une boutique : compléter avec --boutique


JOBS = (
    SheetJob(
        sheet_names=("Multiple list", "Multiple List", "Multiples"),
        spec=TableSpec(
            name="Multiple list",
            columns={
                "sku": ("SKU", "Product", "Item", "Material", "Code"),
                "description": ("Description", "Product Descr", "Product Description"),
                "type": ("Type", "Product Type"),
                "order_multiple": (
                    "Multiple",
                    "Order Multiple",
                    "Order multiple",
                    "UOM",
                    "Multiple of",
                ),
                "units_per_pallet": (
                    "Units per pallet",
                    "Units/Pallet",
                    "Unit/Pallet",
                    "Pallet",
                    "Qty per pallet",
                ),
            },  # fmt: skip
            required=frozenset({"sku", "order_multiple"}),
        ),
        output="master_multiples.csv",
        out_columns=("sku", "description", "type", "order_multiple", "units_per_pallet"),
    ),
    SheetJob(
        sheet_names=("SKU Conversion", "SKU conversion", "Conversion"),
        spec=TableSpec(
            name="SKU Conversion",
            columns={
                "old_sku": ("Old SKU", "Old", "Old Sku", "OLD"),
                "new_sku": ("New SKU", "New", "New Sku", "NEW"),
                "effective_date": ("Date", "Effective Date", "Conversion Date", "Start Date"),
                "comment": ("Comment", "Comments", "Note"),
            },
            required=frozenset({"old_sku", "new_sku", "effective_date"}),
            sku_fields=frozenset({"old_sku", "new_sku"}),
        ),
        output="master_sku_conversions.csv",
        out_columns=("old_sku", "new_sku", "effective_date", "comment"),
        date_fields=("effective_date",),
    ),
    SheetJob(
        sheet_names=("DC Mapping", "DC mapping", "DC"),
        spec=TableSpec(
            name="DC Mapping",
            columns={
                "boutique": ("Boutique", "BTQ", "Store", "Location"),
                "dc": ("DC", "Distribution Center", "Warehouse"),
            },
            required=frozenset({"boutique", "dc"}),
            sku_fields=frozenset({"boutique"}),
        ),
        output="master_dc_mapping.csv",
        out_columns=("boutique", "dc"),
    ),
    SheetJob(
        sheet_names=("Boutique Portfolio", "Portfolio"),
        spec=TableSpec(
            name="Boutique Portfolio",
            columns={
                "boutique": ("Boutique", "BTQ"),
                "sku": ("SKU", "Product", "Item"),
                "filter_out": ("Filter out", "Filter Out", "Filter"),
            },
            required=frozenset({"sku", "filter_out"}),
        ),
        output="master_boutique_portfolio_{boutique}.csv",
        out_columns=("boutique", "sku", "filter_out"),
        fill_boutique=True,
    ),
    # Schedule : la colonne de notes n'est PAS extraite (peut contenir des noms).
    SheetJob(
        sheet_names=("Schedule",),
        spec=TableSpec(
            name="Schedule",
            columns={
                "boutique": ("Boutique", "BTQ"),
                "boutique_name": ("Boutique Name", "Name", "Store"),
                "deliveries_per_week": ("Deliveries per week", "Deliveries/Week", "# Deliveries"),
                "carrier": ("Carrier", "Transporteur"),
                "order_day": ("Order Day", "Order day"),
                "shipping_day": ("Shipping Day", "Ship Day"),
                "delivery_day": ("Delivery Day", "Delivery day"),
                "cover_days": ("Cover Days", "Cover", "Cover days"),
            },
            required=frozenset({"order_day", "delivery_day"}),
            sku_fields=frozenset(),
        ),
        output="master_schedule.csv",
        out_columns=(
            "boutique",
            "boutique_name",
            "deliveries_per_week",
            "carrier",
            "order_day",
            "shipping_day",
            "delivery_day",
            "cover_days",
        ),  # fmt: skip
    ),
)


def find_sheet(names: list[str], wanted: tuple[str, ...]) -> str | None:
    lower = {n.strip().lower(): n for n in names}
    for w in wanted:
        if w.strip().lower() in lower:
            return lower[w.strip().lower()]
    return None


def extract(path: Path, out_dir: Path, boutique: str | None) -> list[ValidationReport]:
    wb = load_workbook(path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()
    out_dir.mkdir(parents=True, exist_ok=True)
    reports: list[ValidationReport] = []
    for job in JOBS:
        sheet = find_sheet(sheet_names, job.sheet_names)
        if sheet is None:
            rep = ValidationReport(source=job.spec.name)
            rep.add(
                "error",
                "SHEET_NOT_FOUND",
                f"Onglet introuvable (cherché : {job.sheet_names}). Onglets : {sheet_names}",
            )
            reports.append(rep)
            continue
        raw = read_raw(path, path.name, sheet=sheet)
        parsed = parse_table(raw, job.spec, f"{path.name} / {sheet}")
        reports.append(parsed.report)
        if parsed.report.has_errors:
            continue
        records = parsed.records
        for f in job.date_fields:
            dates, _, err = parse_dates([r[f] for r in records])
            if err:
                parsed.report.add("error", "BAD_DATE", f"{f} : {err}")
                break
            for r, d in zip(records, dates, strict=True):
                r[f] = d.isoformat() if d else ""
        if parsed.report.has_errors:
            continue
        if job.fill_boutique and any(not r.get("boutique") for r in records):
            if not boutique:
                msg = "Pas de colonne boutique : relancer avec --boutique"
                parsed.report.add("error", "BOUTIQUE_REQUIRED", msg)
                continue
            for r in records:
                r["boutique"] = r.get("boutique") or boutique
        name = job.output.format(boutique=boutique or "ALL")
        with (out_dir / name).open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(job.out_columns)
            for r in records:
                w.writerow([r.get(c, "") for c in job.out_columns])
        parsed.report.add("info", "WRITTEN", f"{len(records)} lignes écrites dans {name}")
    return reports
