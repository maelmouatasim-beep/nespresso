"""Événements SKU : collage de SKU, classeur Excel (template, export, import).

Règles : les SKU restent du texte ; une ligne incomplète est SIGNALÉE avec son numéro
et la raison, jamais ignorée en silence ; la Bible n'est jamais lue (le planner recopie
lui-même ses lignes dans le template).
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from pydantic import ValidationError

from app.domain.dates import MONTHS as MONTHS_FR
from app.domain.events import (
    ALL,
    EFFECT_LABELS,
    FORM_STATUS_LABELS,
    SCOPE_LABELS,
    TYPE_LABELS,
    EventEffect,
    EventQty,
    EventScope,
    EventSku,
    EventStatus,
    EventType,
    SkuEvent,
    default_lead_days,
)
from app.engines.portfolio import numeric_key
from app.ingestion.tables import RawTable, read_raw

# --- Collage de SKU ---------------------------------------------------------------------

_SPLIT = re.compile(r"[,;\s]+")


def split_skus(text: str) -> list[str]:
    """SKU collés depuis Excel : séparés par virgule, point-virgule, espace, tabulation
    ou retour à la ligne. Gardés en texte, dans l'ordre, doublons compris."""
    return [s for s in _SPLIT.split(text or "") if s]


@dataclass(frozen=True)
class SkuCheck:
    sku: str
    description: str | None
    known: bool
    duplicate: bool


def build_catalog(*sources: Iterable[tuple[str, str | None]]) -> dict[str, str]:
    """SKU → description (Multiple list puis Stock Situation)."""
    out: dict[str, str] = {}
    for src in sources:
        for sku, desc in src:
            if not out.get(sku):  # première description non vide
                out[sku] = desc or ""
    return out


def check_skus(skus: Iterable[str], catalog: Mapping[str, str]) -> list[SkuCheck]:
    """Pour chaque SKU : description, connu ou inconnu, doublon (comparaison comme Excel)."""
    by_key = {numeric_key(s): s for s in catalog if numeric_key(s) is not None}
    seen: set[str] = set()
    out = []
    for sku in skus:
        key = numeric_key(sku) or sku
        match = sku if sku in catalog else by_key.get(numeric_key(sku))
        out.append(SkuCheck(sku, catalog.get(match) if match else None, match is not None,
                            key in seen))  # fmt: skip
        seen.add(key)
    return out


# --- Classeur Excel -----------------------------------------------------------------------

EVENTS_SHEET, QTY_SHEET, SKU_SHEET, README_SHEET = (
    "Événements", "Quantités", "SKU détail", "Lisez-moi",
)  # fmt: skip
EVENT_COLUMNS = [
    "Réf", "Portée", "Type", "Titre", "Description", "SKU", "Boutiques", "DC", "Début", "Fin",
    "Anticipation (j)", "Effet", "Qty (toutes boutiques)", "Fichier source", "Propriétaire",
    "Planner responsable", "Statut", "Demandé par", "Date de la demande", "Échéance",
]  # fmt: skip
QTY_COLUMNS = ["Réf", "Boutique", "SKU", "Qty", "Min", "Max", "Déjà envoyé", "Vagues (%)",
               "Vague en cours"]  # fmt: skip
QTY_EXTRA = QTY_COLUMNS[4:]  # facultatives : allocation (vagues) et stock cible (min / max)
SKU_COLUMNS = ["Réf", "SKU", "Multiple", "Nouveau SKU"]
NEXT_ORDER = "Prochaine commande"
README = [
    ("Comment remplir", ""),
    ("Une ligne par événement", "Onglet « Événements ». La colonne Réf relie les onglets."),
    ("Portée", " / ".join(SCOPE_LABELS.values())),
    ("Type", " / ".join(TYPE_LABELS.values())),
    ("Effet", " / ".join(EFFECT_LABELS.values())),
    ("Statut", "Selon les dates (vide) / Brouillon / Annulé. À venir, actif ou terminé se "
               "calcule à partir des dates ; pour terminer, mets une date de fin."),
    ("SKU", "Liste séparée par virgule, point-virgule ou espace. Colonne au format TEXTE "
            "(sinon 7005.70 devient 7005.7)."),
    ("Boutiques", "Vide ou « Toutes » = toutes ; sinon codes séparés (B80, B19)."),
    ("DC", "Facultatif : MT1, TO1, CY1, VA1 (toutes les boutiques du DC)."),
    ("Dates", "Jour/mois/année : 22/10/2026 (ou une cellule date Excel). Début obligatoire ; Fin "
              "facultative (vide = sans fin, ex. une conversion effective à partir d'une date)."),
    ("Qty", "« Qty (toutes boutiques) » ou, plus précis, l'onglet « Quantités » "
            "(Boutique vide ou * = toutes ; SKU vide = tous)."),
    ("Allocation", "Type Promotion, effet « Allocation officielle (vagues) » : Qty = total, "
                   "« Déjà envoyé », « Vagues (%) » (ex. 75/25, vide = une vague) et « Vague en "
                   "cours » dans l'onglet « Quantités ». Remplace "
                   "l'ancien allocations.xlsx."),
    ("Lancement", "Effet « Quantité initiale − Expected » : Qty = quantité initiale, Début = "
                  "date de lancement. Remplace l'ancien launches.xlsx."),
    ("Stock cible", "Type Stock cible : Qty = cible, « Min » et « Max » obligatoires (onglet "
                    "« Quantités »). Sous le min après commande → cible − Expected, arrondie au "
                    "multiple ; au-dessus du max → plafonné. Remplace l'ancien target_stock.xlsx."),
    ("SKU détail", "Facultatif : multiple d'un lancement absent de la Multiple list ; "
                   "« Nouveau SKU » pour une conversion."),
    ("Demande boutique", "Portée Personnel, une seule boutique ; « Demandé par » = prénom ou "
                         "rôle, jamais d'email ; Échéance « Prochaine commande » ou vide."),
    ("Ré-import", "Une Réf « #12 » met à jour l'événement n° 12 ; toute autre Réf en crée un."),
    ("Données", "Exemples fictifs. Ne jamais coller de contacts ni d'emails."),
]  # fmt: skip

_FIELD_NAMES = {"scope": "Portée", "type": "Type", "effect": "Effet", "status": "Statut"}
_LABEL_TO = {
    "scope": {v.lower(): k for k, v in SCOPE_LABELS.items()},
    "type": {v.lower(): k for k, v in TYPE_LABELS.items()},
    "effect": {v.lower(): k for k, v in EFFECT_LABELS.items()},
    "status": {v.lower(): k for k, v in FORM_STATUS_LABELS.items()} | {"actif": EventStatus.ACTIVE},
}


def _codes(text: str) -> tuple[str, ...]:
    codes = tuple(c.strip().upper() for c in re.split(r"[,;\s]+", text or "") if c.strip())
    return () if not codes or codes[0] in ("*", "ALL", "TOUTES") else codes


def _fmt_day(d: date | None) -> str:
    """Jour / mois / année dans le classeur (D-054), relu tel quel à l'import."""
    return d.strftime("%d/%m/%Y") if d else ""


def example_events(business: bool = True) -> list[SkuEvent]:
    """Exemples FICTIFS (template et démo B80 du 05-oct-2026). `business=False` : sans les
    exemples d'allocation, de lancement « quantité initiale » et de stock cible (D-042)."""
    base = date(2026, 10, 5)
    events = [
        SkuEvent(scope=EventScope.TEAM, type=EventType.PROMOTION, title="Festive AOS Oct 14-21",
                 description="Mise en avant des cafés de saison (exemple fictif)",
                 skus=(EventSku(sku="7072.70"), EventSku(sku="7087.70"), EventSku(sku="7094.70")),
                 start=date(2026, 10, 14), end=date(2026, 10, 21), lead_days=7,
                 effect=EventEffect.FLAG, source_file="FESTIVE AOS OCT14-21.xlsx",
                 owner="Marketing (exemple)", planner="Camille"),
        SkuEvent(scope=EventScope.TEAM, type=EventType.LAUNCH, title="Lancement Advent",
                 description="Calendrier de l'Avent, pas encore dans la Multiple list (exemple)",
                 skus=(EventSku(sku="ADV-2026-CA", multiple=6),), dcs=("CY1",),
                 start=date(2026, 10, 12), end=date(2026, 11, 30), lead_days=7,
                 effect=EventEffect.SET_QTY, quantities=(EventQty(qty=20),),
                 source_file="LAUNCH ADVENT 2026.xlsx", planner="Camille"),
        SkuEvent(scope=EventScope.TEAM, type=EventType.STOP, title="Arrêt 7928.70",
                 description="Fin de série : ne plus réapprovisionner (exemple fictif)",
                 skus=(EventSku(sku="7928.70"),), start=base, end=None,
                 effect=EventEffect.BLOCK, planner="Sami"),
        SkuEvent(scope=EventScope.PERSONAL, type=EventType.STORE_REQUEST, title="Demande B80",
                 description="La boutique demande de l'Intenso en plus (exemple fictif)",
                 skus=(EventSku(sku="7042.70"),), boutiques=("B80",), start=base, end=None,
                 effect=EventEffect.ADD_QTY, quantities=(EventQty(qty=240),),
                 planner="Camille", requested_by="Gérance B80", requested_on=base,
                 until_next_order=True),
    ]  # fmt: skip
    if not business:
        return events
    return [
        *events,
        # Ex-onglets Allocations, Lancements, Stocks cibles (D-042), mêmes valeurs qu'avant.
        SkuEvent(scope=EventScope.TEAM, type=EventType.PROMOTION, title="Allocation AOS 14-21 oct",
                 description="Allocation officielle en deux vagues (exemple fictif)",
                 skus=(EventSku(sku="7005.70"),), boutiques=("B80",), start=base, end=None,
                 effect=EventEffect.ALLOCATION,
                 quantities=(EventQty(boutique="B80", sku="7005.70", qty=4800, waves=(75, 25),
                                      current_wave=1),),
                 owner="Business (exemple)", planner="Camille"),
        SkuEvent(scope=EventScope.TEAM, type=EventType.LAUNCH, title="Lancement Festive",
                 description="Quantité initiale avant le lancement (exemple fictif)",
                 skus=(EventSku(sku="7200.70"),), boutiques=("B80",), start=date(2026, 10, 14),
                 end=date(2026, 10, 14), lead_days=9, effect=EventEffect.INITIAL_QTY,
                 quantities=(EventQty(boutique="B80", sku="7200.70", qty=3600),),
                 planner="Camille"),
        SkuEvent(scope=EventScope.TEAM, type=EventType.TARGET_STOCK, title="Stock cible 142172",
                 description="Min / cible / max en boutique (exemple fictif)",
                 skus=(EventSku(sku="142172"),), boutiques=("B80",), start=base, end=None,
                 effect=EventEffect.TARGET,
                 quantities=(EventQty(boutique="B80", sku="142172", qty=30, min_qty=20,
                                      max_qty=60),),
                 planner="Camille"),
        # D-054 : actif du 6 au 22 oct. 2026 (le statut suit les dates, jamais « Terminé » avant).
        SkuEvent(scope=EventScope.TEAM, type=EventType.PROMOTION, title="Free Gift (FG)",
                 description="Pochette offerte à l'achat (exemple fictif)",
                 skus=(EventSku(sku="100924"),), start=date(2026, 10, 6),
                 end=date(2026, 10, 22), effect=EventEffect.FLAG, owner="Marketing (exemple)",
                 planner="Camille"),
    ]  # fmt: skip


DEMO_PLANNERS = ("Camille", "Léa", "Sami")  # prénoms fictifs


def _qty_extra(q: EventQty) -> list[Any]:
    """Min, Max, Déjà envoyé, Vagues (%), Vague en cours ("" = non saisi)."""
    return [
        "" if q.min_qty is None else q.min_qty,
        "" if q.max_qty is None else q.max_qty,
        q.already_sent or "",
        "/".join(map(str, q.waves)),
        q.current_wave if q.waves else "",
    ]


def _waves(text: str) -> tuple[int, ...]:
    parts = [p for p in re.split(r"[/;,+\s]+", text or "") if p]
    out = []
    for part in parts:
        value = _int(part.rstrip("%"), "Vague")
        if value is not None:
            out.append(value)
    return tuple(out)


def events_workbook(events: Iterable[SkuEvent]) -> bytes:
    """Classeur au format du template (export de la page ou template rempli)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = EVENTS_SHEET
    qty_ws, sku_ws, readme = (wb.create_sheet(n) for n in (QTY_SHEET, SKU_SHEET, README_SHEET))
    for sheet, cols in ((ws, EVENT_COLUMNS), (qty_ws, QTY_COLUMNS), (sku_ws, SKU_COLUMNS)):
        sheet.append(cols)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FAF7F2")
            cell.fill = PatternFill("solid", fgColor="3D2B1F")
        sheet.freeze_panes = "A2"
    for i, e in enumerate(events, start=1):
        ref = f"#{e.id}" if e.id else f"E{i}"
        single = [q for q in e.quantities if q.boutique == ALL and q.sku == ALL
                  and not any(v != "" for v in _qty_extra(q))]  # fmt: skip
        ws.append([
            ref, SCOPE_LABELS[e.scope], TYPE_LABELS[e.type], e.title, e.description,
            ", ".join(e.sku_codes()), ", ".join(e.boutiques) or ("" if e.dcs else "Toutes"),
            ", ".join(e.dcs), _fmt_day(e.start), _fmt_day(e.end), e.lead_days,
            EFFECT_LABELS[e.effect], single[0].qty if single else None, e.source_file, e.owner,
            e.planner, FORM_STATUS_LABELS.get(e.status, ""), e.requested_by,
            _fmt_day(e.requested_on),
            NEXT_ORDER if e.until_next_order else "",
        ])  # fmt: skip
        for q in e.quantities:
            extra = _qty_extra(q)
            if not (q.boutique == ALL and q.sku == ALL) or any(v != "" for v in extra):
                qty_ws.append([ref, q.boutique, "" if q.sku == ALL else q.sku, q.qty, *extra])
        for s in e.skus:
            if s.multiple or s.new_sku:
                sku_ws.append([ref, s.sku, s.multiple, s.new_sku or ""])
    for sheet in (ws, qty_ws, sku_ws):  # colonnes en texte : les SKU gardent zéros et décimales
        for col in sheet.iter_cols(min_row=1, max_row=1):
            dim = sheet.column_dimensions[col[0].column_letter]
            dim.width, dim.number_format = 18, "@"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.number_format = "@"
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["F"].width = 30
    for row in README:
        readme.append(row)
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 20
    readme.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Import -------------------------------------------------------------------------------


@dataclass
class ImportResult:
    events: list[SkuEvent] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)  # lignes incomplètes, avec la raison
    warnings: list[str] = field(default_factory=list)


def _records(raw: RawTable, columns: list[str]) -> list[tuple[int, dict[str, str], set[str]]]:
    """(n° de ligne, valeurs, colonnes saisies comme nombre) ; en-tête cherché, pas supposé."""
    norm = [c.lower() for c in columns]
    for h, row in enumerate(raw.rows[:10]):
        cells = [str(c).strip().lower() for c in row]
        if norm[0] in cells and sum(c in cells for c in norm) >= min(3, len(norm)):
            idx = {c: cells.index(c.lower()) for c in columns if c.lower() in cells}
            out = []
            for r in range(h + 1, len(raw.rows)):
                row_r = raw.rows[r]
                if not any(str(c).strip() for c in row_r):
                    continue
                vals = {c: (row_r[i].strip() if i < len(row_r) else "") for c, i in idx.items()}
                numeric = {c for c, i in idx.items() if (r, i) in raw.numeric_cells}
                out.append((r + 1, vals, numeric))
            return out
    raise ValueError(f"En-têtes introuvables (attendu : {', '.join(columns[:4])}…)")


_FR_MONTHS = {m.rstrip("."): i for i, m in enumerate(MONTHS_FR, start=1)} | {
    "janvier": 1, "fevr": 2, "février": 2, "fev": 2, "fév": 2, "avril": 4, "juillet": 7,
    "aout": 8, "septembre": 9, "sep": 9, "octobre": 10, "novembre": 11, "décembre": 12,
    "decembre": 12, "dec": 12,
}  # fmt: skip


def _day(text: str) -> date | None:
    """Date du classeur : cellule date Excel, JJ/MM/AAAA (jour d'abord, jamais mois d'abord),
    AAAA-MM-JJ ou « 22 oct. 2026 ». Année sur 2 chiffres refusée (ambiguë)."""
    text = text.strip()
    if not text:
        return None
    iso = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?", text)
    dmy = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", text)
    named = re.fullmatch(r"(\d{1,2})[\s-]+([^\W\d_]+)\.?[\s-]+(\d{2,4})", text)
    try:
        if iso:
            return date(int(iso[1]), int(iso[2]), int(iso[3]))
        if dmy or named:
            m = dmy or named
            assert m is not None
            if len(m[3]) != 4:
                raise ValueError(f"date « {text} » : année sur 4 chiffres (ex. 22/10/2026)")
            month = int(m[2]) if dmy else _FR_MONTHS.get(m[2].lower())
            if month is None:
                raise ValueError(f"date « {text} » : mois inconnu")
            return date(int(m[3]), month, int(m[1]))
    except ValueError as exc:
        if "date «" in str(exc):
            raise
        raise ValueError(f"date impossible « {text} » (jour/mois/année, ex. 22/10/2026)") from exc
    raise ValueError(f"date illisible « {text} » (jour/mois/année, ex. 22/10/2026)")


def _label(kind: str, text: str, default: Any = None) -> Any:
    if not text:
        if default is None:
            raise ValueError(f"{_FIELD_NAMES[kind]} manquant")
        return default
    value = _LABEL_TO[kind].get(text.strip().lower())
    if kind == "status" and text.strip().lower() == "terminé":
        raise ValueError("« Terminé » n'est plus un statut à saisir : mets une date de fin "
                         "(le statut suit les dates)")  # fmt: skip
    if value is None:
        raise ValueError(f"valeur inconnue « {text} »")
    return value


def _int(text: str, what: str) -> int | None:
    if not text:
        return None
    try:
        value = float(text.replace(",", ".").replace(" ", "").replace(" ", ""))
    except ValueError as exc:
        raise ValueError(f"{what} illisible « {text} »") from exc
    if not value.is_integer() or value < 0:
        raise ValueError(f"{what} doit être un entier positif (« {text} »)")
    return int(value)


def read_events_workbook(data: bytes, filename: str, author: str) -> ImportResult:
    """Lit un classeur au format du template. Les lignes incomplètes sont listées."""
    result = ImportResult()
    try:
        main = _records(read_raw(io.BytesIO(data), filename, EVENTS_SHEET), EVENT_COLUMNS)
    except (ValueError, KeyError) as exc:
        result.problems.append(f"Onglet « {EVENTS_SHEET} » : {exc}")
        return result
    extras: dict[str, list[tuple[int, dict[str, str], set[str]]]] = {}
    for sheet, cols in ((QTY_SHEET, QTY_COLUMNS), (SKU_SHEET, SKU_COLUMNS)):
        try:
            extras[sheet] = _records(read_raw(io.BytesIO(data), filename, sheet), cols)
        except (ValueError, KeyError):
            extras[sheet] = []  # onglet facultatif
    for line, rec, numeric in main:
        where = f"{EVENTS_SHEET}, ligne {line}"
        ref = rec.get("Réf", "")
        try:
            if "SKU" in numeric:
                result.warnings.append(f"{where} : SKU saisi comme nombre, vérifier les zéros "
                                       "et décimales (formater la colonne en Texte)")  # fmt: skip
            event = _event_from(rec, ref, extras, author)
            result.events.append(event)
        except (ValueError, ValidationError) as exc:
            msg = "; ".join(e["msg"].removeprefix("Value error, ") for e in exc.errors()) \
                if isinstance(exc, ValidationError) else str(exc)  # fmt: skip
            result.problems.append(f"{where} ({rec.get('Titre') or 'sans titre'}) : {msg}")
    return result


def _event_from(rec: dict[str, str], ref: str, extras: dict[str, Any], author: str) -> SkuEvent:
    event_type = _label("type", rec.get("Type", ""))
    scope = _label("scope", rec.get("Portée", ""))
    effect = _label("effect", rec.get("Effet", ""))
    title = rec.get("Titre", "")
    if not title:
        raise ValueError("titre manquant")
    skus = list(dict.fromkeys(split_skus(rec.get("SKU", ""))))
    if not skus:
        raise ValueError("aucun SKU")
    start = _day(rec.get("Début", ""))
    if start is None:
        raise ValueError("date de début manquante")
    details = {r["SKU"]: r for _, r, _ in extras.get(SKU_SHEET, []) if r.get("Réf") == ref and ref}
    items = []
    for s in skus:
        old, _, new = s.partition(">")
        d = details.get(old, {})
        items.append(EventSku(sku=old, multiple=_int(d.get("Multiple", ""), "Multiple"),
                              new_sku=(new or d.get("Nouveau SKU") or None)))  # fmt: skip
    quantities = []
    single = _int(rec.get("Qty (toutes boutiques)", ""), "Qty")
    if single is not None:
        quantities.append(EventQty(qty=single))
    for _, q, _ in extras.get(QTY_SHEET, []):
        if ref and q.get("Réf") == ref:
            value = _int(q.get("Qty", ""), "Qty")
            if value is None:
                raise ValueError(f"quantité vide dans l'onglet {QTY_SHEET}")
            b = q.get("Boutique", "").upper()
            quantities.append(EventQty(
                boutique=ALL if b in ("", "*", "TOUTES") else b, sku=q.get("SKU") or ALL,
                qty=value, min_qty=_int(q.get("Min", ""), "Min"),
                max_qty=_int(q.get("Max", ""), "Max"),
                already_sent=_int(q.get("Déjà envoyé", ""), "Déjà envoyé") or 0,
                waves=_waves(q.get("Vagues (%)", "")),
                current_wave=_int(q.get("Vague en cours", ""), "Vague en cours") or 1,
            ))  # fmt: skip
    lead = _int(rec.get("Anticipation (j)", ""), "Anticipation")
    existing = int(ref[1:]) if re.fullmatch(r"#\d+", ref) else None
    return SkuEvent(
        id=existing, scope=scope, type=event_type, title=title,
        description=rec.get("Description", ""), skus=tuple(items),
        boutiques=_codes(rec.get("Boutiques", "")), dcs=_codes(rec.get("DC", "")),
        start=start, end=_day(rec.get("Fin", "")),
        lead_days=default_lead_days(event_type) if lead is None else lead,
        effect=effect, quantities=tuple(quantities), source_file=rec.get("Fichier source", ""),
        owner=rec.get("Propriétaire", ""), planner=rec.get("Planner responsable", ""),
        status=_label("status", rec.get("Statut", ""), EventStatus.ACTIVE),
        requested_by=rec.get("Demandé par", ""),
        requested_on=_day(rec.get("Date de la demande", "")),
        until_next_order=rec.get("Échéance", "").strip().lower() == NEXT_ORDER.lower(),
        created_by=author,
    )  # fmt: skip
