"""Lecture générique de tableaux (CSV ou Excel) avec détection d'en-têtes et rapport.

Principes (CLAUDE.md, règles 2, 3 et 9) :
- tout est lu en TEXTE ; un SKU stocké comme nombre dans Excel est signalé ;
- la ligne d'en-tête est cherchée, pas supposée : on accepte plusieurs noms possibles
  par colonne (alias) et on explique ce qui manque ;
- aucune valeur n'est remplacée en silence : chaque problème devient une entrée du
  rapport de validation.
"""

from __future__ import annotations

import csv
import io
import re
import warnings
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO, Literal

try:  # lecteur Excel rapide ; openpyxl reste utilisé s'il n'est pas installé
    from python_calamine import CalamineWorkbook
except ImportError:  # pragma: no cover
    CalamineWorkbook = None

MAX_HEADER_SCAN = 25
FOOTER_PATTERNS = ("no filters applied", "filters applied", "applied filters")


# --- Rapport de validation ---------------------------------------------------------


@dataclass(frozen=True)
class ValidationIssue:
    level: Literal["error", "warning", "info"]
    code: str
    message: str
    rows: tuple[int, ...] = ()


@dataclass
class ValidationReport:
    source: str
    issues: list[ValidationIssue] = field(default_factory=list)
    rows_read: int = 0
    rows_kept: int = 0

    def add(self, level: Literal["error", "warning", "info"], code: str, message: str,
            rows: Sequence[int] = ()) -> None:  # fmt: skip
        self.issues.append(ValidationIssue(level, code, message, tuple(rows[:20])))

    @property
    def has_errors(self) -> bool:
        return any(i.level == "error" for i in self.issues)

    def as_rows(self) -> list[dict[str, str]]:
        return [
            {
                "fichier": self.source,
                "niveau": i.level.upper(),
                "code": i.code,
                "message": i.message,
                "lignes (exemples)": ", ".join(map(str, i.rows)),
            }
            for i in self.issues
        ]


# --- Description d'un tableau attendu ----------------------------------------------


@dataclass(frozen=True)
class TableSpec:
    name: str
    columns: dict[str, tuple[str, ...]]  # champ -> noms d'en-tête acceptés
    required: frozenset[str]
    sku_fields: frozenset[str] = frozenset({"sku"})


def _norm(text: object) -> str:
    return re.sub(r"[\s_]+", " ", str(text or "")).strip().lower()


@dataclass(frozen=True)
class RawTable:
    """Cellules en texte + positions des SKU qui étaient des nombres dans Excel."""

    rows: list[list[str]]
    numeric_cells: set[tuple[int, int]]  # (index de ligne, index de colonne)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_excel_calamine(
    source: Path | BinaryIO, sheet: str | None, max_rows: int | None
) -> RawTable:
    """Lecture rapide (≈ 1 s pour 90 000 lignes au lieu de ≈ 15 s avec openpyxl)."""
    assert CalamineWorkbook is not None
    if isinstance(source, Path):
        wb = CalamineWorkbook.from_path(str(source))
    else:
        source.seek(0)
        wb = CalamineWorkbook.from_filelike(source)
    try:
        ws = wb.get_sheet_by_name(sheet) if sheet else wb.get_sheet_by_index(0)
        data = ws.to_python(skip_empty_area=False, nrows=max_rows)
    finally:
        wb.close()
    rows: list[list[str]] = []
    numeric: set[tuple[int, int]] = set()
    for r, row in enumerate(data):
        values = []
        for c, value in enumerate(row):
            if isinstance(value, int | float) and not isinstance(value, bool):
                numeric.add((r, c))
            values.append(_cell_text(value))
        rows.append(values)
    return RawTable(rows=rows, numeric_cells=numeric)


def read_raw(
    source: Path | BinaryIO, filename: str, sheet: str | None = None, max_rows: int | None = None
) -> RawTable:
    """Lit un .csv / .xlsx / .xlsm en cellules texte (`max_rows` : seulement le début)."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"} and CalamineWorkbook is not None:
        return _read_excel_calamine(source, sheet, max_rows)
    if suffix in {".xlsx", ".xlsm"}:
        with warnings.catch_warnings():
            # Les exports Power BI n'ont pas de style par défaut : avertissement sans intérêt.
            warnings.filterwarnings("ignore", message="Workbook contains no default style")
            from openpyxl import load_workbook

            wb = load_workbook(source, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows: list[list[str]] = []
            numeric: set[tuple[int, int]] = set()
            for r, row in enumerate(ws.iter_rows(values_only=True, max_row=max_rows)):
                values = []
                for c, value in enumerate(row):
                    if isinstance(value, int | float) and not isinstance(value, bool):
                        numeric.add((r, c))
                    values.append(_cell_text(value))
                rows.append(values)
            return RawTable(rows=rows, numeric_cells=numeric)
        finally:
            wb.close()
    if suffix in {".csv", ".txt"}:
        data = source.read_bytes() if isinstance(source, Path) else source.read()
        text = data.decode("utf-8-sig", errors="strict") if isinstance(data, bytes) else data
        sample = text[:5000]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows_csv = [[c.strip() for c in row] for row in reader]
        return RawTable(rows=rows_csv[:max_rows] if max_rows else rows_csv, numeric_cells=set())
    raise ValueError(f"{filename} : format non pris en charge ({suffix}). Utiliser .csv ou .xlsx")


@dataclass(frozen=True)
class ParsedTable:
    records: list[dict[str, str]]  # champ -> texte ; "_row" = numéro de ligne dans le fichier
    report: ValidationReport


def parse_table(raw: RawTable, spec: TableSpec, source: str) -> ParsedTable:
    """Trouve l'en-tête, mappe les colonnes, retire les lignes de pied et vides."""
    report = ValidationReport(source=source)
    aliases = {f: {_norm(a) for a in names} for f, names in spec.columns.items()}
    header_idx, mapping = None, {}
    for i, row in enumerate(raw.rows[:MAX_HEADER_SCAN]):
        cand: dict[str, int] = {}
        for c, cell in enumerate(row):
            for f, names in aliases.items():
                if _norm(cell) in names and f not in cand:
                    cand[f] = c
        if spec.required <= cand.keys():
            header_idx, mapping = i, cand
            break
    if header_idx is None:
        seen = [r for r in raw.rows[:5] if any(r)]
        expected = {f: spec.columns[f][0] for f in sorted(spec.required)}
        report.add(
            "error",
            "HEADER_NOT_FOUND",
            f"En-têtes introuvables pour « {spec.name} ». Colonnes attendues : "
            f"{list(expected.values())}. Premières lignes lues : {seen[:3]}",
        )
        return ParsedTable(records=[], report=report)

    missing_optional = sorted(set(spec.columns) - mapping.keys())
    if missing_optional:
        report.add(
            "info",
            "OPTIONAL_COLUMNS_MISSING",
            f"Colonnes facultatives absentes : {missing_optional}",
        )

    records: list[dict[str, str]] = []
    footer_rows, empty_rows, numeric_sku_rows = [], [], []
    for r in range(header_idx + 1, len(raw.rows)):
        row = raw.rows[r]
        line_no = r + 1
        if not any(cell for cell in row):
            continue
        if any(_norm(cell).startswith(FOOTER_PATTERNS) for cell in row if cell):
            footer_rows.append(line_no)
            continue
        rec = {f: (row[c] if c < len(row) else "") for f, c in mapping.items()}
        rec["_row"] = str(line_no)
        if any(not rec.get(f) for f in spec.sku_fields & mapping.keys()):
            empty_rows.append(line_no)
            continue
        if any((r, mapping[f]) in raw.numeric_cells for f in spec.sku_fields & mapping.keys()):
            numeric_sku_rows.append(line_no)
        records.append(rec)
    report.rows_read = len(raw.rows) - header_idx - 1
    if footer_rows:
        report.add(
            "info", "FOOTER_REMOVED", "Ligne de pied « No filters applied » retirée", footer_rows
        )
    if empty_rows:
        report.add(
            "warning", "EMPTY_SKU", f"{len(empty_rows)} ligne(s) sans SKU ignorée(s)", empty_rows
        )
    if numeric_sku_rows:
        report.add(
            "warning",
            "SKU_AS_NUMBER",
            f"{len(numeric_sku_rows)} SKU stocké(s) comme nombre dans Excel : zéros initiaux ou "
            "décimales peut-être perdus (ex. 7005.70 → 7005.7). Formater la colonne en Texte.",
            numeric_sku_rows,
        )
    report.rows_kept = len(records)
    return ParsedTable(records=records, report=report)


def parse_number(text: str) -> float | None:
    """Nombre au format « 1234.5 », « 1 234,5 » ou « 1,234.5 ». None si illisible."""
    t = text.replace(" ", "").replace(" ", "")
    if not t:
        return None
    if "," in t and "." in t:
        t = (
            t.replace(",", "")
            if t.rfind(".") > t.rfind(",")
            else t.replace(".", "").replace(",", ".")
        )
    elif "," in t:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


DATE_FORMATS = ("%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%m-%d-%Y")


def parse_dates(values: Sequence[str]) -> tuple[list[date | None], str | None, str | None]:
    """Convertit une colonne de dates. Retourne (dates, format retenu, erreur).

    ISO (2026-09-28) est accepté directement. Pour les autres formats, la colonne
    entière doit être lisible avec UN SEUL format ; si deux formats marchent
    (ex. 03/04/2026), c'est ambigu et on refuse plutôt que de deviner.
    """

    def iso(v: str) -> date | None:
        try:
            return datetime.fromisoformat(v).date()
        except ValueError:
            return None

    parsed = [iso(v) for v in values]
    if all(p is not None for p in parsed):
        return parsed, "ISO", None
    working = []
    for fmt in DATE_FORMATS:
        try:
            working.append((fmt, [datetime.strptime(v.split(" ")[0], fmt).date() for v in values]))
        except ValueError:
            continue
    if len(working) == 1:
        return list(working[0][1]), working[0][0], None
    if not working:
        bad = [v for v, p in zip(values, parsed, strict=True) if p is None][:5]
        return [None] * len(values), None, f"Dates illisibles, exemples : {bad}"
    distinct = {tuple(d) for _, d in working}
    if len(distinct) == 1:
        return list(working[0][1]), working[0][0], None
    return (
        [None] * len(values),
        None,
        (
            f"Format de date ambigu (jour/mois ou mois/jour ?) : formats possibles "
            f"{[f for f, _ in working]}. Exporter les dates au format AAAA-MM-JJ."
        ),
    )
