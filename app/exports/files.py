"""Fichiers produits en mémoire (octets) pour le téléchargement : classeur Excel détaillé
d'une commande et zip des fichiers LT d'une journée. Aucune écriture sur disque."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable, Mapping
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.models import OVERRIDE_CATEGORY_LABELS, OverrideCategory
from app.engines.export import lt_content, lt_filename
from app.services.planning import PlanningResult

HEADER_FILL = PatternFill("solid", fgColor="EFE6DC")
STATUS_FILL = {"REVIEW": "FBEED0", "BLOCKED": "F8E0E0"}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: Iterable[list], text_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for idx in text_cols:  # SKU en texte : jamais converti en nombre par Excel
        for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
            for c in cell:
                c.number_format = "@"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(60, len(h) + 4))
    ws.freeze_panes = "A2"
    return ws


def order_workbook(
    result: PlanningResult,
    *,
    author: str = "",
    sources: Mapping[str, str] | None = None,
    overrides: Iterable = (),
) -> bytes:
    """Classeur de la commande : lignes + explications, paramètres, overrides, palettes."""
    wb = Workbook()
    wb.remove(wb.active)
    p = result.params
    rows = []
    for x in result.lines:
        e, inp = x.explanation, x.explanation["inputs"]
        rows.append(
            [
                x.sku,
                x.description or "",
                x.product_type or "",
                x.status.value,
                x.source.value,
                x.forecast_qty if x.forecast_qty is not None else x.qty,
                x.qty,
                inp["sales_total"],
                inp["expected_total"],
                e["current_cover_days"],
                e["post_cover_days"],
                e["cover_applied"],
                e["raw_need"],
                e["multiple"],
                e["multiple_source"],
                inp["dc_available"],
                "oui" if x.qty_reliable else "NON",
                " | ".join(x.reasons),
            ]  # fmt: skip
        )
    ws = _sheet(
        wb,
        "Commande",
        [
            "SKU",
            "Description",
            "Type",
            "Statut",
            "Origine",
            "Qty prévision",
            "Qty finale",
            f"Ventes {result.history_days} j",
            "Expected",
            "Couv. actuelle (j)",
            "Couv. après (j)",
            "Cover appliqué",
            "Besoin brut",
            "Multiple",
            "Source du multiple",
            "Stock DC",
            "Qty fiable",
            "Signalements",
        ],  # fmt: skip
        rows,
        text_cols=(1,),
    )
    for r in range(2, ws.max_row + 1):
        fill = STATUS_FILL.get(str(ws.cell(r, 4).value))
        if fill:
            ws.cell(r, 4).fill = PatternFill("solid", fgColor=fill)
    s = result.summary
    _sheet(
        wb,
        "Paramètres",
        ["Élément", "Valeur"],
        [
            ["Boutique", p.boutique],
            ["DC", result.dc],
            ["Date du run", p.run_date.isoformat()],
            ["Date de livraison", p.delivery_date.isoformat()],
            ["Jours de couverture", p.cover_days],
            ["Cover café", p.coffee_cover_days if p.coffee_cover_days is not None else "—"],
            ["Mode", result.mode.value],
            ["Historique de ventes (jours)", result.history_days],
            ["Planner", author],
            ["Lignes commandées", s.lines_ordered],
            ["Unités", s.units_total],
            ["Palettes estimées", round(result.pallets.total_pallets, 2)],
            ["REVIEW", s.by_status["REVIEW"]],
            ["BLOCKED", s.by_status["BLOCKED"]],
            *[[f"Version {k}", v] for k, v in result.rule_versions.items()],
            *[[f"Source {k}", v] for k, v in (sources or {}).items()],
            *[["Alerte", w] for w in result.rule_warnings],
        ],  # fmt: skip
    )
    _sheet(
        wb,
        "Overrides",
        ["SKU", "Avant", "Après", "Catégorie", "Raison", "Auteur", "Horodatage"],
        [
            [
                o.sku,
                o.qty_before,
                o.qty_after,
                OVERRIDE_CATEGORY_LABELS[OverrideCategory(o.category)],
                o.reason,
                o.author,
                o.timestamp.isoformat(timespec="minutes"),
            ]  # fmt: skip
            for o in overrides
        ],
        text_cols=(1,),
    )
    _sheet(
        wb,
        "Palettes",
        ["Famille", "Unités", "Palettes"],
        [[f, u, round(result.pallets.pallets[f], 2)] for f, u in result.pallets.units.items()],
    )
    _sheet(wb, "Lignes exclues", ["SKU", "Raison"], sorted(result.excluded.items()), text_cols=(1,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lt_zip(results: Iterable[PlanningResult], when: datetime) -> tuple[str, bytes]:
    """Zip des fichiers LT d'une journée (une boutique = un fichier). Retourne (nom, octets)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for r in results:
            if any(x.qty > 0 for x in r.lines):
                z.writestr(lt_filename(r.params.boutique, when), lt_content(r.lines))
    return f"LT {when:%Y-%m-%d %H_%M}.zip", buf.getvalue()
