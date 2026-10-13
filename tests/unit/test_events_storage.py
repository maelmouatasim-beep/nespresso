"""Événements SKU : base locale (historique, planners, validations) et classeur Excel."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.domain.events import EventEffect, EventScope, EventSku, EventStatus, EventType, SkuEvent
from app.ingestion.events import (
    build_catalog,
    check_skus,
    events_workbook,
    example_events,
    read_events_workbook,
    split_skus,
)
from app.storage import db
from app.storage import events as store

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "evenements_sku.xlsx"


@pytest.fixture
def conn():  # noqa: ANN201
    c = db.connect(Path(":memory:"))
    yield c
    c.close()


def test_paste_keeps_skus_as_text_and_flags_unknown_and_duplicates() -> None:
    skus = split_skus("7005.70, 0005452;3517/BULK\n7005.7\tXYZ 9999.70")
    assert skus == ["7005.70", "0005452", "3517/BULK", "7005.7", "XYZ", "9999.70"]
    catalog = build_catalog([("7005.70", "VER-Melozio"), ("0005452", "Sac"), ("3517/BULK", None)])
    checks = {c.sku: c for c in check_skus(skus, catalog)}
    assert checks["7005.70"].description == "VER-Melozio" and checks["7005.70"].known
    assert checks["7005.7"].duplicate and checks["7005.7"].known  # 7005.7 = 7005.70
    assert not checks["9999.70"].known and not checks["XYZ"].known
    assert checks["0005452"].sku == "0005452"


def test_planners_first_name_only(conn) -> None:  # noqa: ANN001
    store.add_planner(conn, "  Camille ")
    store.add_planner(conn, "camille")  # même prénom, pas de doublon
    store.add_planner(conn, "Sami")
    assert store.list_planners(conn) == ["Camille", "Sami"]
    with pytest.raises(ValueError, match="pas d'email"):
        store.add_planner(conn, "camille@example.com")
    store.remove_planner(conn, "Sami")
    assert store.list_planners(conn) == ["Camille"]


def test_event_history_keeps_who_changed_what(conn) -> None:  # noqa: ANN001
    e = example_events()[0]
    saved = store.save_event(conn, e, "Camille")
    assert saved.id == 1 and saved.created_by == "Camille" and saved.created_at
    store.save_event(conn, saved.model_copy(update={"title": "Festive AOS v2"}), "Sami")
    store.set_status(conn, 1, EventStatus.DONE, "Sami")
    store.log_takeover(conn, 1, "Léa")
    hist = store.event_history(conn, 1)
    assert [h["action"] for h in hist] == ["repris", "terminé", "modifié", "créé"]
    assert hist[2]["author"] == "Sami" and hist[2]["changes"]["title"] == [
        "Festive AOS Oct 14-21", "Festive AOS v2"]  # fmt: skip
    again = store.get_event(conn, 1)
    assert again is not None and again.status is EventStatus.DONE
    assert again.created_by == "Camille" and again.skus[0].sku == "7072.70"
    with pytest.raises(ValueError, match="Choisis ton nom"):
        store.save_event(conn, e, " ")


def test_validations_are_traced(conn) -> None:  # noqa: ANN001
    store.record_validation(conn, author="Camille", boutique="B80", delivery=date(2026, 10, 8),
                            kind="confirm_block", skus=["7921.70"])  # fmt: skip
    rows = store.list_validations(conn, "B80")
    assert rows[0]["author"] == "Camille" and rows[0]["skus"] == ["7921.70"] and rows[0]["at"]


def test_template_exists_and_roundtrips_with_fictional_examples() -> None:
    data = TEMPLATE.read_bytes()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Événements", "Quantités", "SKU détail", "Lisez-moi"]
    res = read_events_workbook(data, TEMPLATE.name, "Camille")
    assert not res.problems, res.problems
    assert [e.title for e in res.events] == [e.title for e in example_events()]
    advent = res.events[1]
    assert advent.skus[0].multiple == 6 and advent.dcs == ("CY1",) and advent.lead_days == 7
    assert res.events[3].until_next_order and res.events[3].scope is EventScope.PERSONAL


def test_import_reports_incomplete_rows_instead_of_ignoring_them() -> None:
    from openpyxl import Workbook

    good = SkuEvent(scope=EventScope.TEAM, type=EventType.BAN, title="OK", effect=EventEffect.BLOCK,
                    skus=(EventSku(sku="0005452"),), start=date(2026, 10, 5))  # fmt: skip
    wb = load_workbook(io.BytesIO(events_workbook([good])))
    ws = wb["Événements"]
    ws.append(["E2", "Équipe", "Promotion", "Sans fin", "", "7005.70", "", "", "2026-10-05", "",
               0, "Flag seulement"])  # fmt: skip
    ws.append(["E3", "Équipe", "Promotion", "Sans SKU", "", "", "", "", "2026-10-05",
               "2026-10-09", 0, "Flag seulement"])  # fmt: skip
    ws.append(["E4", "Équipe", "Type bizarre", "X", "", "1", "", "", "2026-10-05"])
    buf = io.BytesIO()
    wb.save(buf)
    res = read_events_workbook(buf.getvalue(), "e.xlsx", "Camille")
    assert [e.title for e in res.events] == ["OK", "Sans fin"]  # fin vide permise (D-053)
    assert res.events[0].skus[0].sku == "0005452" and res.events[1].end is None
    assert len(res.problems) == 2
    assert "aucun SKU" in res.problems[0] and "Type bizarre" in res.problems[1]
    del Workbook
