"""Tests : lecteurs Power BI, modèles métier, extraction du .xlsm, SQLite, service complet."""

from __future__ import annotations

import io
import json
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.domain.models import ForecastMode, Override, PlanningParameters
from app.ingestion.business_inputs import read_allocations, read_launches, read_targets
from app.ingestion.fixtures_loader import load_fixture_dir
from app.ingestion.masters import build_inputs_for_boutique, load_masters_dir
from app.ingestion.powerbi import check_freshness, read_movements, read_stock_situation
from app.ingestion.tables import parse_dates, parse_number
from app.services.planning import BusinessInputs, run_planning
from app.storage import db

ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "tests" / "fixtures" / "b80_2026-10-05"
sys.path.insert(0, str(ROOT))


def csv_bytes(text: str) -> io.BytesIO:
    return io.BytesIO(text.encode("utf-8"))


def xlsx_bytes(rows: list[list[object]]) -> io.BytesIO:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


STOCK_HEADER = (
    "Stock,Product,Product Descr,Product Type,Expected,Available,Incoming,Reserved,Waiting"
)


# --- Stock Situation ----------------------------------------------------------------------


def test_stock_situation_footer_negative_and_empty_type() -> None:
    text = (
        f"{STOCK_HEADER}\n"
        "B80,0005452,Cups,,5,5,0,0,0\n"
        "B80,7005.70,VER-Melozio,C,-10,0,0,0,10\n"
        "No filters applied,,,,,,,,\n"
    )
    lines, rep = read_stock_situation(csv_bytes(text), "export.csv")
    assert [x.sku for x in lines] == ["0005452", "7005.70"]
    codes = {i.code for i in rep.issues}
    assert {"FOOTER_REMOVED", "NEGATIVE_EXPECTED", "EMPTY_PRODUCT_TYPE"} <= codes
    assert not rep.has_errors


def test_stock_situation_header_not_on_first_row_and_semicolons() -> None:
    text = "Export Power BI;;\n\nStock;Product;Expected;Available\nB80;7005.70;1 234,5;1234,5\n"
    lines, rep = read_stock_situation(csv_bytes(text), "export.csv")
    assert lines[0].expected == 1234.5
    assert any(i.code == "STOCK_COLUMNS_MISSING" for i in rep.issues)


def test_stock_situation_duplicates_and_bad_numbers_are_errors() -> None:
    text = f"{STOCK_HEADER}\nB80,A,x,C,1,1,0,0,0\nB80,A,x,C,1,1,0,0,0\nB80,B,x,C,abc,1,0,0,0\n"
    _, rep = read_stock_situation(csv_bytes(text), "export.csv")
    assert {"DUPLICATE_SKU", "BAD_NUMBER"} <= {i.code for i in rep.issues if i.level == "error"}


def test_missing_header_is_explained() -> None:
    _, rep = read_stock_situation(csv_bytes("a,b\n1,2\n"), "x.csv")
    assert rep.has_errors and "Colonnes attendues" in rep.issues[0].message


def test_xlsx_numeric_sku_is_flagged() -> None:
    buf = xlsx_bytes([["Stock", "Product", "Expected", "Available"], ["B80", 7005.7, 1, 1],
                      ["B80", "0005452", 2, 2]])  # fmt: skip
    lines, rep = read_stock_situation(buf, "export.xlsx")
    assert [x.sku for x in lines] == ["7005.7", "0005452"]
    assert any(i.code == "SKU_AS_NUMBER" and i.rows == (2,) for i in rep.issues)


def test_fixture_stock_situation_reads_cleanly() -> None:
    lines, rep = read_stock_situation(FIXTURE / "stock_situation_B80_CY1.csv", "ss.csv")
    assert len(lines) == 3723
    assert not rep.has_errors


# --- Mouvements et dates --------------------------------------------------------------------


def test_movements_aliases_and_ambiguous_dates() -> None:
    ok = "Stock,Product,Date,Movement Code,Quantity\nB80,7005.70,9/28/2026,8730,2\n"
    moves, rep = read_movements(csv_bytes(ok), "m.csv")
    assert moves[0].movement_date == date(2026, 9, 28)
    assert any(i.code == "DATE_FORMAT" for i in rep.issues)
    amb = "Stock,Product,Date,Movement Code,Quantity\nB80,A,03/04/2026,8730,2\n"
    moves, rep = read_movements(csv_bytes(amb), "m.csv")
    assert moves == [] and any(i.code == "BAD_DATE" for i in rep.issues)


def test_parse_helpers() -> None:
    assert parse_number("1,234.5") == 1234.5
    assert parse_number("1 234,5") == 1234.5
    assert parse_number("x") is None
    assert parse_dates(["2026-09-28"])[0] == [date(2026, 9, 28)]


def test_freshness_warnings() -> None:
    moves, rep = read_movements(FIXTURE / "stock_movements_B80.csv", "m.csv")
    check_freshness(rep, date(2026, 10, 10), movements=moves, extracted_on=date(2026, 10, 9))
    assert {"STALE_SALES", "STALE_EXTRACT"} <= {i.code for i in rep.issues}


# --- Modèles métier -----------------------------------------------------------------------------


def test_business_templates_read() -> None:
    a, rep = read_allocations(csv_bytes(
        "boutique,sku,total_qty,already_sent,wave_plan,current_wave\n"
        "B80,7005.70,1000,0,75/25,1\nB80,X,10,0,50/20,1\n"), "a.csv")  # fmt: skip
    assert len(a) == 1 and a[0].wave_plan == (75, 25)
    assert rep.has_errors  # 50/20 ne fait pas 100
    lz, rep = read_launches(
        csv_bytes("boutique,sku,launch_date,qty\nB80,7200.70,2026-10-14,3600\n"), "l.csv"
    )
    assert lz[0].launch_date == date(2026, 10, 14) and not rep.has_errors
    t, rep = read_targets(
        csv_bytes("boutique,sku,min_qty,target_qty,max_qty\nB80,A,1,2,3\n"), "t.csv"
    )
    assert t[0].target_qty == 2 and not rep.has_errors


def test_generated_templates_are_readable(tmp_path: Path) -> None:
    for name, reader in (("allocations.xlsx", read_allocations), ("launches.xlsx", read_launches),
                         ("target_stock.xlsx", read_targets)):  # fmt: skip
        items, rep = reader(ROOT / "templates" / name, name)
        assert items == [] and not rep.has_errors, rep.as_rows()


# --- Extraction du calculateur -----------------------------------------------------------------


def test_extract_masters_from_workbook(tmp_path: Path) -> None:
    from tools.extract_masters import extract

    wb = Workbook()
    ws = wb.active
    ws.title = "Multiple list"
    ws.append(["Calculateur B80"])
    ws.append(["SKU", "Description", "Type", "Multiple", "Units per pallet"])
    ws.append(["7005.70", "VER-Melozio", "C", 240, 8640])
    ws.append(["0005452", "Cups", "LC", 12, None])
    conv = wb.create_sheet("SKU Conversion")
    conv.append(["Old SKU", "New SKU", "Date", "Comment"])
    conv.append(["7039.70", "7005.70", datetime(2026, 9, 2), ""])
    dc = wb.create_sheet("DC Mapping")
    dc.append(["Boutique", "DC"])
    dc.append(["B80", "CY1"])
    pf = wb.create_sheet("Boutique Portfolio")
    pf.append(["SKU", "Filter out"])
    pf.append(["7005.70", "No"])
    sc = wb.create_sheet("Schedule")
    sc.append(["Boutique", "Order Day", "Delivery Day", "Cover Days", "Notes"])
    sc.append(["B80", "Monday", "Wednesday", 9, "Appeler Jean"])
    path = tmp_path / "B80 105.xlsm"
    wb.save(path)

    reports = extract(path, tmp_path / "out", "B80")
    assert not any(r.has_errors for r in reports), [r.as_rows() for r in reports]
    masters = load_masters_dir(tmp_path / "out")
    assert masters.products["0005452"].order_multiple == 12
    assert masters.conversions[0].effective_date == date(2026, 9, 2)
    assert masters.portfolio["B80"]["7005.70"] == "No"
    schedule = (tmp_path / "out" / "master_schedule.csv").read_text()
    assert "Jean" not in schedule  # les notes (données personnelles) ne sont pas extraites


def test_extract_masters_reports_missing_sheet(tmp_path: Path) -> None:
    from tools.extract_masters import extract

    wb = Workbook()
    wb.active.title = "Autre"
    path = tmp_path / "x.xlsx"
    wb.save(path)
    reports = extract(path, tmp_path / "out", "B80")
    assert all(r.has_errors for r in reports)
    assert "Onglets" in reports[0].issues[0].message


# --- Service complet + SQLite --------------------------------------------------------------------


@pytest.fixture(scope="module")
def b80_run():
    data = load_fixture_dir(FIXTURE)
    inputs, dc = build_inputs_for_boutique(
        "B80", data.stock_situation, data.movements, data.products, data.conversions,
        data.dc_mapping,
    )  # fmt: skip
    params = PlanningParameters(boutique="B80", run_date=date(2026, 10, 5),
                                delivery_date=date(2026, 10, 7), cover_days=9,
                                coffee_cover_days=None)  # fmt: skip
    return inputs, dc, params, data.portfolio


def test_run_planning_end_to_end(b80_run) -> None:
    inputs, dc, params, portfolio = b80_run
    res = run_planning(inputs, params, ForecastMode.SAFE, dc=dc, portfolio=portfolio)
    by_sku = {x.sku: x for x in res.lines}
    assert by_sku["7005.70"].qty == 3120
    assert "2007.70" in by_sku  # Q-002 : retenu par la règle (40 ventes)
    assert res.summary.lines_ordered > 0
    assert res.pallets.total_pallets > 0


def test_run_planning_with_golden_skus_matches_parity(b80_run) -> None:
    import pandas as pd

    inputs, dc, params, portfolio = b80_run
    g = pd.read_csv(FIXTURE / "golden_stock_cover_final_B80.csv", dtype=str)
    res = run_planning(inputs, params, ForecastMode.EXCEL_PARITY, dc=dc, portfolio=portfolio,
                       skus=list(g["sku"]))  # fmt: skip
    qty = {x.sku: x.qty for x in res.lines}
    formula = g[g["excel_cell_is_formula"] == "True"]
    assert all(
        qty[s] == int(q)
        for s, q in zip(formula["sku"], formula["excel_suggested_qty"], strict=True)
    )


def test_storage_roundtrip(b80_run) -> None:
    inputs, dc, params, portfolio = b80_run
    override = Override(sku="7005.70", qty_before=3120, qty_after=3600, reason="Promo AOS",
                        author="Planner A", timestamp=datetime(2026, 10, 5, 11, 51))  # fmt: skip
    res = run_planning(inputs, params, ForecastMode.EXCEL_PARITY, dc=dc, portfolio=portfolio,
                       business_inputs=BusinessInputs(), overrides=[override])  # fmt: skip
    conn = db.connect(Path(":memory:"))
    first = db.save_run(conn, res, author="Planner A", sources={"ss": "abc"}, overrides=[override])
    second = db.save_run(conn, res, author="Planner B", sources={"ss": "def"})
    db.mark_exported(conn, second, "B80 05-Oct-2026-11_51.csv")
    runs = db.list_runs(conn, "B80")
    assert [r["id"] for r in runs] == [second, first]
    assert runs[0]["lt_filename"] == "B80 05-Oct-2026-11_51.csv"
    assert db.run_quantities(conn, first)["7005.70"] == 3600
    assert db.run_overrides(conn, first)[0]["reason"] == "Promo AOS"
    assert db.previous_run_id(conn, "B80", before_id=second) == first
    stored = conn.execute("SELECT explanation FROM recommendations WHERE sku='7005.70'").fetchone()
    assert json.loads(stored[0])["override"]["author"] == "Planner A"
    with pytest.raises(ValueError):
        db.save_run(conn, res, author=" ", sources={})


def test_masters_dir_from_fixture() -> None:
    m = load_masters_dir(FIXTURE)
    assert m.dc_mapping["B80"] == "CY1"
    assert "B80" in m.portfolio
