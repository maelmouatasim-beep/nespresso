"""Tests : Schedule, règles boutique, journée multi-boutiques, exports Excel / zip,
catégories d'override et migration de la base."""

from __future__ import annotations

import io
import sqlite3
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.domain.models import (
    BoutiqueRule,
    ForecastMode,
    Override,
    PlanningParameters,
    ScheduleSlot,
    Status,
)
from app.engines.pallets import estimate_pallets
from app.engines.rules import check_boutique_rules
from app.engines.schedule import boutiques_ordering_on, next_delivery_date, parse_weekday
from app.exports.files import lt_zip, order_workbook
from app.ingestion.business_inputs import read_rules
from app.ingestion.fixtures_loader import load_fixture_dir
from app.ingestion.masters import build_inputs_for_boutique, load_masters_dir
from app.services.planning import BusinessInputs, run_day, run_planning, with_overrides
from app.storage import db
from tests.unit.test_business_and_outputs import line

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "b80_2026-10-05"
TUESDAY = date(2026, 10, 6)


# --- Schedule ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "day"),
    [("Monday", 0), ("TUESDAY", 1), ("wed", 2), ("jeudi", 3), ("Ven", 4), ("", None), ("x", None)],
)
def test_parse_weekday(text: str, day: int | None) -> None:
    assert parse_weekday(text) == day


def test_next_delivery_date_is_strictly_after_run() -> None:
    assert next_delivery_date(TUESDAY, 3) == date(2026, 10, 8)  # jeudi
    assert next_delivery_date(TUESDAY, 1) == date(2026, 10, 13)  # mardi suivant
    assert next_delivery_date(TUESDAY, 0) == date(2026, 10, 12)  # lundi suivant


def test_boutiques_ordering_on_fixture_schedule() -> None:
    masters = load_masters_dir(FIXTURE)
    assert masters.schedule and not masters.schedule_issues
    day = {d.boutique: d.delivery_date for d in boutiques_ordering_on(masters.schedule, TUESDAY)}
    assert day["B80"] == date(2026, 10, 8)  # B80 : commande mardi, livraison jeudi
    assert "B1" not in day  # B1 commande lundi, mercredi, jeudi


def test_schedule_keeps_first_slot_per_boutique() -> None:
    slots = [ScheduleSlot(boutique="B1", order_weekday=1, delivery_weekday=3),
             ScheduleSlot(boutique="B1", order_weekday=1, delivery_weekday=4)]  # fmt: skip
    assert len(boutiques_ordering_on(slots, TUESDAY)) == 1


# --- Règles boutique ---------------------------------------------------------------------------


def test_rule_needs_sku_only_for_sku_rules() -> None:
    with pytest.raises(ValidationError):
        BoutiqueRule(boutique="B80", rule="max_qty_sku", value=10)
    with pytest.raises(ValidationError):
        BoutiqueRule(boutique="B80", rule="max_pallets", sku="7005.70", value=4)


def test_rules_flag_without_changing_quantities() -> None:
    lines = [line("7005.70", 3120, desc="VER-Melozio"), line("142172", 0, ptype="LC", desc="Mug")]
    rules = [
        BoutiqueRule(boutique="B80", rule="max_qty_sku", sku="7005.70", value=2400),
        BoutiqueRule(boutique="B80", rule="min_qty_sku", sku="142172", value=6),
        BoutiqueRule(boutique="B80", rule="max_pallets", value=0.1, comment="max 4 pallets"),
        BoutiqueRule(boutique="B80", rule="max_units", value=100),
        BoutiqueRule(boutique="B1", rule="max_units", value=1),  # autre boutique : ignorée
    ]
    pallets = estimate_pallets(lines, {}, "CY1")
    out, warnings = check_boutique_rules(lines, rules, "B80", pallets)
    assert [x.qty for x in out] == [3120, 0]
    assert all(x.status is Status.REVIEW for x in out)
    assert len(warnings) == 2 and "max 4 pallets" in warnings[0]


def test_rules_reader_and_template() -> None:
    text = "boutique,rule,sku,value,comment\nB80,max_pallets,,4,note\nB80,max_qty_sku,,10,\n"
    rules, rep = read_rules(io.BytesIO(text.encode()), "r.csv")
    assert len(rules) == 1 and rep.has_errors  # 2e ligne : SKU manquant
    root = Path(__file__).resolve().parents[2]
    rules, rep = read_rules(root / "templates" / "boutique_rules.xlsx", "boutique_rules.xlsx")
    assert rules == [] and not rep.has_errors


# --- Service : règles conservées après overrides, journée ------------------------------------


@pytest.fixture(scope="module")
def b80():
    data = load_fixture_dir(FIXTURE)
    inputs, dc = build_inputs_for_boutique(
        "B80", data.stock_situation, data.movements, data.products, data.conversions,
        data.dc_mapping,
    )  # fmt: skip
    params = PlanningParameters(boutique="B80", run_date=date(2026, 10, 5),
                                delivery_date=date(2026, 10, 8), cover_days=9,
                                coffee_cover_days=None)  # fmt: skip
    return data, inputs, dc, params


def test_rules_reapplied_once_after_overrides(b80) -> None:
    data, inputs, dc, params = b80
    rule = BoutiqueRule(boutique="B80", rule="max_qty_sku", sku="7005.70", value=3000)
    res = run_planning(inputs, params, ForecastMode.SAFE, dc=dc, portfolio=data.portfolio,
                       business_inputs=BusinessInputs(rules=(rule,)))  # fmt: skip
    melo = next(x for x in res.lines if x.sku == "7005.70")
    assert melo.status is Status.REVIEW
    ov = Override(sku="7005.70", qty_before=3120, qty_after=2880, category="overstock",
                  reason="Place limitée", author="A", timestamp=datetime(2026, 10, 5))  # fmt: skip
    final = with_overrides(res, [ov], inputs.products)
    melo = next(x for x in final.lines if x.sku == "7005.70")
    assert melo.qty == 2880
    assert not any(r.startswith("Règle boutique") for r in melo.reasons)  # 2880 ≤ 3000
    again = with_overrides(final, [], inputs.products)
    melo = next(x for x in again.lines if x.sku == "7005.70")
    assert melo.qty == 3120 and sum(r.startswith("Règle boutique") for r in melo.reasons) == 1


def test_run_day_isolates_errors(b80) -> None:
    data, inputs, dc, params = b80

    def build(boutique: str):
        return build_inputs_for_boutique(
            boutique, data.stock_situation, data.movements, data.products, data.conversions,
            data.dc_mapping,
        )  # fmt: skip

    other = params.model_copy(update={"boutique": "B1"})
    out = run_day([params, other], ForecastMode.SAFE, build_inputs=build,
                  portfolio_by_boutique={"B80": data.portfolio})  # fmt: skip
    assert out[0].result is not None and out[0].error is None
    assert out[1].result is None and "B1" in (out[1].error or "")


# --- Exports ----------------------------------------------------------------------------------


def test_order_workbook_and_zip(b80) -> None:
    data, inputs, dc, params = b80
    when = datetime(2026, 10, 5, 11, 51)
    ov = Override(sku="7005.70", qty_before=3120, qty_after=3600, category="promo",
                  reason="Promo AOS", author="A", timestamp=when)  # fmt: skip
    res = run_planning(inputs, params, ForecastMode.SAFE, dc=dc, portfolio=data.portfolio,
                       overrides=[ov])  # fmt: skip
    wb = load_workbook(io.BytesIO(order_workbook(res, author="A", overrides=[ov])))
    assert wb.sheetnames == ["Commande", "Paramètres", "Overrides", "Palettes", "Lignes exclues"]
    rows = {r[0]: r for r in wb["Commande"].iter_rows(min_row=2, values_only=True)}
    assert rows["7005.70"][6] == 3600 and rows["7005.70"][5] == 3120
    assert isinstance(next(iter(rows)), str)
    assert wb["Overrides"]["D2"].value == "Promo / événement"

    name, payload = lt_zip([res], datetime(2026, 10, 5, 11, 51))
    with zipfile.ZipFile(io.BytesIO(payload)) as z:
        assert z.namelist() == ["B80 05-Oct-2026-11_51.csv"]
        assert "7005.70;3600" in z.read(z.namelist()[0]).decode()
    assert name.endswith(".zip")


# --- Base : catégories et migration -----------------------------------------------------------


def test_override_category_required() -> None:
    with pytest.raises(ValidationError):
        Override(sku="A", qty_before=1, qty_after=2, reason="xyz", author="A",
                 timestamp=datetime(2026, 1, 1))  # type: ignore[call-arg]  # fmt: skip


def test_storage_migrates_old_overrides_table(tmp_path: Path, b80) -> None:
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript(
        """CREATE TABLE overrides (id INTEGER PRIMARY KEY AUTOINCREMENT, run_id INTEGER NOT NULL,
           sku TEXT NOT NULL, qty_before INTEGER NOT NULL, qty_after INTEGER NOT NULL,
           reason TEXT NOT NULL, author TEXT NOT NULL, timestamp TEXT NOT NULL);"""
    )
    old.close()
    data, inputs, dc, params = b80
    ov = Override(sku="7005.70", qty_before=3120, qty_after=3600, category="launch",
                  reason="Lancement", author="A", timestamp=datetime(2026, 10, 5))  # fmt: skip
    res = run_planning(inputs, params, ForecastMode.SAFE, dc=dc, portfolio=data.portfolio,
                       overrides=[ov])  # fmt: skip
    conn = db.connect(path)
    run_id = db.save_run(conn, res, author="A", sources={}, overrides=[ov])
    assert db.run_overrides(conn, run_id)[0]["category"] == "launch"
    stats = db.override_stats(conn, "B80")
    assert stats[0]["sku"] == "7005.70" and stats[0]["boutique"] == "B80"
