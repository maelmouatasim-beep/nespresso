"""Format réel des exports Power BI (vérifié le 07-oct-2026 sur les fichiers du planner).

Aucune donnée réelle ici : les classeurs sont fabriqués avec les MÊMES en-têtes, la même
ligne de pied et les mêmes conventions (sorties en négatif, SKU avec espace en trop).
"""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from app.ingestion.depot import MOVES_KIND, STOCK_KIND, detect_extraction_date, detect_kind
from app.ingestion.powerbi import read_movements, read_stock_situation

STOCK_HEADER = ("Stock", "Product", "Product Descr", "Product Type", "Expected", "Available",
                "Incoming", "Reserved", "Waiting")  # fmt: skip
MOVES_HEADER = ("Stock Movement Id", "Stock", "Product Nr", "Product Type (Prod)", "Stock Mvt Date",
                "Mvt Code", "Quantity (Sum)", "Mvt Code Descr (Mvt Cd)")  # fmt: skip
MOVES_FOOTER = ("Applied filters:\nRelative Date (absolute, max) is greater than or equal to -7\n"
                "Quantity (Sum) is less than or equal to 0")  # fmt: skip


def xlsx(header: tuple, rows: list[tuple], footer: str = "No filters applied") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(header)
    for r in rows:
        ws.append(r)
    ws.append(())
    ws.append((footer,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stock_file(rows: list[tuple] | None = None) -> bytes:
    return xlsx(STOCK_HEADER, rows or [
        ("B80", "7005.70", "VER-Melozio", "C", 7800, 7800, 0, 0, 0),
        ("B80", "10081", "Recipe Card", "D", 0, 0, 0, 0, 0),
        ("CY1", "7005.70", "VER-Melozio", "C", 368080, 368080, 0, 0, 0),
    ])  # fmt: skip


def moves_file(quantities: tuple[int, ...] = (-20, -330, -1)) -> bytes:
    day = datetime(2026, 10, 6)
    rows = [
        ("251238506", "B80", "7005.70", "C", day, "2801", quantities[0], "CA-LAST CHANCE BOGO"),
        ("251238507", "B80", "7005.70", "C", day, "8730", quantities[1], "CA-BIM Order B2C"),
        ("251238508", "B80", "10081", "D", day, "1005", quantities[2], "CA-Internet Order"),
    ]
    return xlsx(MOVES_HEADER, rows, MOVES_FOOTER)


def test_movements_real_headers_negative_exits_become_positive_sales() -> None:
    moves, report = read_movements(io.BytesIO(moves_file()), "data - 2026-10-07T084848.694.xlsx")
    assert not report.has_errors
    assert [m.quantity for m in moves] == [20, 330, 1]
    assert moves[0].movement_id == "251238506" and moves[0].movement_code == "2801"
    assert moves[0].movement_date == date(2026, 10, 6)
    assert moves[0].movement_description == "CA-LAST CHANCE BOGO"
    codes = {i.code for i in report.issues}
    assert {"NEGATIVE_EXITS", "FOOTER_REMOVED"} <= codes


def test_movements_positive_quantities_are_kept_as_is() -> None:
    moves, report = read_movements(io.BytesIO(moves_file((20, 330, 1))), "m.xlsx")
    assert [m.quantity for m in moves] == [20, 330, 1]
    assert "NEGATIVE_EXITS" not in {i.code for i in report.issues}


def test_movements_mixed_signs_are_refused() -> None:
    moves, report = read_movements(io.BytesIO(moves_file((-20, 330, -1))), "m.xlsx")
    assert moves == []
    assert "MIXED_SIGNS" in {i.code for i in report.issues if i.level == "error"}


def test_stock_situation_real_export_reads() -> None:
    lines, report = read_stock_situation(io.BytesIO(stock_file()), "s.xlsx")
    assert not report.has_errors
    assert {(x.location, x.sku, x.expected) for x in lines} >= {("B80", "7005.70", 7800)}


def test_stock_duplicate_with_a_space_keeps_the_line_with_stock() -> None:
    rows = [
        ("B5", " 473ECO/B", "ORI-Essenza Auto C101 Black", "M", 0, 0, 0, 0, 0),
        ("B5", "473ECO/B", "ORI-US Essenza Auto C101 US Black", "M", 3, 3, 0, 0, 0),
    ]
    lines, report = read_stock_situation(io.BytesIO(stock_file(rows)), "s.xlsx")
    assert not report.has_errors
    assert [(x.sku, x.expected) for x in lines] == [("473ECO/B", 3)]
    assert "DUPLICATE_SKU_MERGED" in {i.code for i in report.issues}


def test_stock_duplicate_with_stock_on_both_lines_is_an_error() -> None:
    rows = [
        ("B5", " 473ECO/B", "A", "M", 2, 2, 0, 0, 0),
        ("B5", "473ECO/B", "B", "M", 3, 3, 0, 0, 0),
    ]
    _, report = read_stock_situation(io.BytesIO(stock_file(rows)), "s.xlsx")
    assert "DUPLICATE_SKU" in {i.code for i in report.issues if i.level == "error"}


def test_detect_kind_by_columns_whatever_the_file_name() -> None:
    assert detect_kind("data - 2026-10-07T084852.664.xlsx", stock_file()) == STOCK_KIND
    assert detect_kind("data - 2026-10-07T084848.694.xlsx", moves_file()) == MOVES_KIND
    assert detect_kind("autre.xlsx", xlsx(("A", "B"), [(1, 2)])) is None
    assert detect_kind("pas_excel.xlsx", b"not a workbook") is None


def test_power_bi_file_name_gives_the_extraction_date() -> None:
    extraction = detect_extraction_date("data - 2026-10-07T084852.664.xlsx", stock_file())
    assert extraction.day == date(2026, 10, 7)


def test_workspace_deposit_load_and_cache(tmp_path: Path) -> None:
    from app.ui.workspace import MOVES, STOCK, Workspace

    ws = Workspace(tmp_path)
    for kind, data in ((STOCK, stock_file()), (MOVES, moves_file())):
        name = f"data - 2026-10-07T0848{kind}.xlsx"
        ws.save_depot(kind, name, data, detect_extraction_date(name, data), date(2026, 10, 7))
    ws.use_test_referentials("Planner")
    assert ws.has_masters()
    first = ws.load_day()
    assert len(list(ws.depot_dir.glob("*/_cache_*.pkl"))) == 2
    second = ws.load_day()  # depuis le cache
    assert second.stock == first.stock and second.movements == first.movements
    assert [m.quantity for m in second.movements] == [20, 330, 1]
    assert not first.has_errors


def test_import_calculator_writes_the_referentials(tmp_path: Path) -> None:
    from app.ui.workspace import Workspace

    wb = Workbook()
    ml = wb.active
    ml.title = "Multiple list"
    ml.append(("SKU", "Description", "Type", "Multiple", "Units per pallet"))
    ml.append(("7005.70", "VER-Melozio", "C", 240, 8000))
    conv = wb.create_sheet("SKU Conversion")
    conv.append(("Old SKU", "New SKU", "Date", "Comment"))
    conv.append(("7039.70", "7005.70", datetime(2026, 9, 2), ""))
    dc = wb.create_sheet("DC Mapping")
    dc.append(("Boutique", "DC"))
    dc.append(("B80", "CY1"))
    buf = io.BytesIO()
    wb.save(buf)
    ws = Workspace(tmp_path)
    reports = ws.import_calculator("B80 105.xlsx", buf.getvalue(), "B80", "Planner")
    assert ws.has_masters()
    assert ws.meta()["multiples"]["updated_by"] == "Planner"
    assert any(i.code == "SHEET_NOT_FOUND" for r in reports for i in r.issues)  # pas de Schedule
    assert not (tmp_path / "_tmp" / "B80 105.xlsx").exists()
