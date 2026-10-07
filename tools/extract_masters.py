"""Extrait les tables maîtres d'un calculateur Excel (ex. « B80 105.xlsm ») en CSV.

À lancer EN LOCAL uniquement : le fichier Excel et les CSV produits restent dans
`data/` (jamais dans le dépôt).

    python tools/extract_masters.py "data/B80 105.xlsm" --boutique B80 --out data/masters
    python tools/extract_masters.py "data/B80 105.xlsm" --inspect   # voir les onglets / en-têtes

La structure n'est pas supposée : pour chaque onglet, l'en-tête est cherché parmi
plusieurs noms possibles. Si un en-tête n'est pas trouvé, le script s'arrête et
affiche les premières lignes de l'onglet pour qu'on ajoute le bon nom.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.ingestion.calculator import extract  # noqa: E402


def inspect(path: Path, rows: int = 8) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            print(f"\n=== Onglet « {ws.title} » ({ws.sheet_state})")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= rows:
                    break
                print("  ", [c for c in row[:15]])
    finally:
        wb.close()


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--out", type=Path, default=ROOT / "data" / "masters")
    ap.add_argument("--boutique", help="Code boutique du calculateur (ex. B80)")
    ap.add_argument("--inspect", action="store_true", help="Afficher onglets et premières lignes")
    args = ap.parse_args()
    if args.inspect:
        inspect(args.workbook)
        return 0
    reports = extract(args.workbook, args.out, args.boutique)
    errors = False
    for rep in reports:
        for row in rep.as_rows():
            print(f"[{row['niveau']}] {row['fichier']} : {row['message']}")
        errors |= rep.has_errors
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
