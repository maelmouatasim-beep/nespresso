"""Génère les modèles Excel vierges à remplir par l'équipe (dossier templates/).

python tools/make_templates.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates"
sys.path.insert(0, str(ROOT))

from app.engines.categories import CATEGORIES, DEFAULT_CATEGORY_MAP  # noqa: E402

TEMPLATES = {
    "allocations.xlsx": (
        ["boutique", "sku", "total_qty", "already_sent", "wave_plan", "current_wave", "comment"],
        [
            "boutique : code boutique (ex. B80).",
            "sku : SKU en TEXTE (formater la colonne en Texte avant de coller).",
            "total_qty : quantité totale allouée à la boutique.",
            "already_sent : quantité déjà envoyée lors des vagues précédentes (0 sinon).",
            "wave_plan : pourcentages des vagues séparés par /, total 100 (ex. 100/0 ou 75/25).",
            "current_wave : numéro de la vague à envoyer maintenant (1, 2…).",
            "Quantité envoyée = total × % cumulé jusqu'à la vague en cours − déjà envoyé.",
            "Une allocation REMPLACE la prévision pour ce SKU (elle n'est jamais écrasée).",
        ],
    ),
    "launches.xlsx": (
        ["boutique", "sku", "launch_date", "qty", "comment"],
        [
            "launch_date : date de lancement (format AAAA-MM-JJ ou date Excel).",
            "qty : quantité initiale à avoir en boutique au lancement.",
            "Jusqu'à la date de lancement : commande = qty − Expected, arrondi au multiple.",
            "Après la date : la prévision normale reprend.",
        ],
    ),
    "target_stock.xlsx": (
        ["boutique", "sku", "min_qty", "target_qty", "max_qty"],
        [
            "Stock après commande (Expected + commande) comparé à min / cible / max.",
            "En dessous du min : on remonte à la cible. Au-dessus du max : on plafonne.",
            "Ne s'applique pas aux allocations ni aux lancements.",
        ],
    ),
    "exclusions.xlsx": (
        ["sku", "boutiques", "reason", "author", "updated_on"],
        [
            "SKU à ne jamais envoyer (ex. SKU e-commerce non vendu en boutique).",
            "boutiques : ALL pour toutes, sinon codes séparés par ; (ex. B80;B1).",
            "reason : pourquoi ce SKU est exclu. author : qui l'a décidé.",
            "updated_on : date de la décision (AAAA-MM-JJ).",
            "L'exclusion apparaît dans le « Pourquoi » de la ligne (quantité 0, BLOCKED).",
        ],
    ),
    "categories.xlsx": (
        ["product_type", "categorie"],
        [
            "Classement des types produit dans le filtre « Catégorie » de la commande.",
            "categorie = Cafés, Machines, Accessoires, Consommables & sacs ou Autres.",
            "Un type absent de la table est classé « Autres ».",
            "Affichage seulement : la catégorie ne change aucune quantité.",
            "Le modèle est prérempli avec le classement validé par le planner.",
        ],
    ),
    "boutique_rules.xlsx": (
        ["boutique", "rule", "sku", "value", "comment"],
        [
            "Règles boutique tirées des notes en texte libre, validées par un humain.",
            "rule = max_pallets (palettes max pour la commande, sans SKU)",
            "       max_units (unités max pour la commande, sans SKU)",
            "       max_qty_sku (quantité max pour un SKU, SKU obligatoire)",
            "       min_qty_sku (quantité min pour un SKU, SKU obligatoire)",
            "Une règle ne change jamais une quantité : elle passe la ligne en REVIEW",
            "ou affiche une alerte sur la commande.",
            "Exemple : B80 | max_pallets | | 4 | « max 4 pallets » (note de la Bible).",
        ],
    ),
}


PREFILLED = {
    "categories.xlsx": sorted(
        DEFAULT_CATEGORY_MAP.items(), key=lambda kv: (CATEGORIES.index(kv[1]), kv[0])
    ),
}


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for name, (headers, notes) in TEMPLATES.items():
        wb = Workbook()
        ws = wb.active
        ws.title = "Données"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EFE6DC")
        for row in PREFILLED.get(name, []):
            ws.append(list(row))
        sku_cols = [chr(65 + i) for i, h in enumerate(headers) if h in ("boutique", "sku")]
        for col in sku_cols:
            for row in range(2, 501):
                ws[f"{col}{row}"].number_format = "@"  # texte : protège les SKU
        for i, _ in enumerate(headers):
            ws.column_dimensions[chr(65 + i)].width = 16
        help_ws = wb.create_sheet("Mode d'emploi")
        for line in notes:
            help_ws.append([line])
        help_ws.column_dimensions["A"].width = 100
        wb.save(OUT / name)
        print(f"Modèle écrit : {OUT / name}")


if __name__ == "__main__":
    main()
